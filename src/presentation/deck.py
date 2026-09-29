from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(13, 23, 38)
NAVY_2 = RGBColor(22, 36, 58)
WHITE = RGBColor(248, 250, 252)
MUTED = RGBColor(148, 163, 184)
GRID = RGBColor(51, 65, 85)
GREEN = RGBColor(34, 197, 94)
RED = RGBColor(239, 68, 68)
BLUE = RGBColor(56, 189, 248)


def _set_bg(slide, color: RGBColor = NAVY) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12.0), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.72), Inches(1.12), Inches(11.8), Inches(0.4))
        sp = s.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Aptos"
        sp.font.size = Pt(11)
        sp.font.color.rgb = MUTED


def _add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12.0), Inches(0.28))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(8)
    p.font.color.rgb = MUTED


def _metric_card(slide, x: float, y: float, w: float, h: float, label: str, value: str, positive: bool | None = None) -> None:
    shape = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY_2
    shape.line.color.rgb = GRID
    tf = shape.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.name = "Aptos"
    p1.font.size = Pt(11)
    p1.font.color.rgb = MUTED
    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.name = "Aptos Display"
    p2.font.size = Pt(21)
    p2.font.bold = True
    p2.font.color.rgb = GREEN if positive is True else RED if positive is False else WHITE


def _add_bullets(slide, bullets: Iterable[str], x: float, y: float, w: float, h: float, font_size: int = 16) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = WHITE
        p.space_after = Pt(9)


def _ranked_bar_chart(slide, items: list[dict[str, Any]], x: float, y: float, w: float, h: float) -> None:
    if not items:
        _add_bullets(slide, ["Not available in this run."], x, y, w, h)
        return

    ordered = sorted(items, key=lambda row: float(row.get("return_pct", 0.0)), reverse=True)
    values = [float(row.get("return_pct", 0.0)) for row in ordered]
    min_v = min(min(values), 0.0)
    max_v = max(max(values), 0.0)
    span = max_v - min_v or 1.0

    label_w = 1.25
    value_w = 0.85
    plot_x = x + label_w
    plot_w = w - label_w - value_w
    zero_x = plot_x + plot_w * ((0.0 - min_v) / span)

    axis = slide.shapes.add_shape(1, Inches(zero_x), Inches(y), Inches(0.012), Inches(h))
    axis.fill.solid(); axis.fill.fore_color.rgb = GRID; axis.line.color.rgb = GRID

    row_h = h / max(len(ordered), 1)
    for i, row in enumerate(ordered):
        symbol = str(row.get("symbol", ""))
        value = float(row.get("return_pct", 0.0))
        cy = y + i * row_h + row_h * 0.18
        bh = row_h * 0.50

        lab = slide.shapes.add_textbox(Inches(x), Inches(cy - 0.02), Inches(label_w - 0.1), Inches(bh + 0.12))
        lp = lab.text_frame.paragraphs[0]
        lp.text = symbol
        lp.font.name = "Aptos"
        lp.font.size = Pt(12)
        lp.font.color.rgb = WHITE
        lp.alignment = PP_ALIGN.RIGHT

        if value >= 0:
            bar_x = zero_x
            bar_w = plot_w * (value / span)
            color = GREEN
        else:
            bar_w = plot_w * (abs(value) / span)
            bar_x = zero_x - bar_w
            color = RED
        bar_w = max(bar_w, 0.02)
        bar = slide.shapes.add_shape(1, Inches(bar_x), Inches(cy), Inches(bar_w), Inches(bh))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.color.rgb = color

        val = slide.shapes.add_textbox(Inches(x + w - value_w + 0.05), Inches(cy - 0.02), Inches(value_w - 0.05), Inches(bh + 0.12))
        vp = val.text_frame.paragraphs[0]
        vp.text = f"{value:+.2f}%"
        vp.font.name = "Aptos"
        vp.font.size = Pt(11)
        vp.font.bold = True
        vp.font.color.rgb = color
        vp.alignment = PP_ALIGN.RIGHT


def _two_column_list(slide, left_title: str, left: list[dict[str, Any]], right_title: str, right: list[dict[str, Any]]) -> None:
    for x, title, items, color in [(0.8, left_title, left, GREEN), (6.9, right_title, right, RED)]:
        head = slide.shapes.add_textbox(Inches(x), Inches(1.65), Inches(5.3), Inches(0.4))
        hp = head.text_frame.paragraphs[0]
        hp.text = title
        hp.font.bold = True
        hp.font.size = Pt(17)
        hp.font.color.rgb = color
        if not items:
            _add_bullets(slide, ["Not available"], x, 2.15, 5.2, 3.8)
            continue
        y = 2.15
        for row in items[:5]:
            card = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(5.2), Inches(0.72))
            card.fill.solid(); card.fill.fore_color.rgb = NAVY_2; card.line.color.rgb = GRID
            tf = card.text_frame; tf.clear()
            p = tf.paragraphs[0]
            p.text = str(row.get("symbol", ""))
            p.font.bold = True; p.font.size = Pt(16); p.font.color.rgb = WHITE
            p2 = tf.add_paragraph()
            move = row.get("return_pct")
            catalyst = row.get("catalyst")
            suffix = f" · {catalyst}" if catalyst else ""
            p2.text = (f"{move:+.2f}%" if isinstance(move, (int, float)) else "Not available") + suffix
            p2.font.size = Pt(10); p2.font.color.rgb = MUTED
            y += 0.82


def build_deck(report: dict[str, Any], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    meta = report.get("meta", {})
    start = meta.get("start_date", "")
    end = meta.get("end_date", "")
    benchmarks = report.get("benchmarks", []) or []
    sources = report.get("sources", []) or []
    footer = "Sources: " + ", ".join(sources[:4]) if sources else "Deterministic market-data inputs"

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, f"Sunday Market Recap | {start} → {end}", report.get("one_line_week", "Weekly U.S. market snapshot"))
    for i, row in enumerate(benchmarks[:5]):
        _metric_card(slide, 0.72 + i * 2.48, 2.15, 2.18, 1.25, str(row.get("symbol")), f"{float(row.get('return_pct', 0)):+.2f}%", float(row.get("return_pct", 0)) >= 0)
    _add_bullets(slide, report.get("highlights", [])[:3], 0.82, 4.0, 11.6, 2.1, 17)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "Major benchmarks", "Weekly return, ranked best to worst")
    _ranked_bar_chart(slide, benchmarks, 0.9, 1.65, 7.3, 4.9)
    _add_bullets(slide, report.get("benchmark_takeaways", [])[:3], 8.55, 1.8, 3.9, 3.5, 15)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, report.get("leadership_title", "Leadership & relative performance"))
    rel = report.get("relative_performance", []) or []
    if rel:
        y = 1.75
        for row in rel[:6]:
            _metric_card(slide, 0.9, y, 3.2, 0.82, str(row.get("label", "Spread")), f"{float(row.get('value_pct', 0)):+.2f} pp", float(row.get("value_pct", 0)) >= 0)
            y += 0.92
    else:
        _add_bullets(slide, ["Relative-performance spreads not available."], 0.9, 1.8, 4.8, 2.0)
    _add_bullets(slide, report.get("leadership_notes", [])[:4], 4.6, 1.8, 7.6, 4.7, 17)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "Top movers", "Weekly movers only; daily moves are labeled separately")
    movers = report.get("movers", {}) or {}
    _two_column_list(slide, "Top gainers", movers.get("gainers", []) or [], "Top losers", movers.get("losers", []) or [])
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    sectors = report.get("sectors", []) or []
    _add_title(slide, "Sector performance")
    if sectors:
        _ranked_bar_chart(slide, sectors, 0.9, 1.65, 8.0, 4.9)
        _add_bullets(slide, report.get("sector_notes", [])[:3], 9.1, 1.8, 3.4, 3.8, 14)
    else:
        _add_bullets(slide, ["Sector data not available in this run."], 1.0, 2.2, 10.8, 2.0, 22)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, report.get("breadth_title", "Breadth / participation"))
    breadth = report.get("breadth", {}) or {}
    metrics = [("Advancers", breadth.get("advancers")), ("Decliners", breadth.get("decliners")), ("Positive week", breadth.get("positive_pct")), ("52w highs", breadth.get("new_highs")), ("52w lows", breadth.get("new_lows"))]
    shown = 0
    for label, value in metrics:
        if value is None:
            continue
        val = f"{value:.0f}%" if label == "Positive week" and isinstance(value, (int, float)) else str(value)
        _metric_card(slide, 0.85 + shown * 2.45, 2.0, 2.1, 1.2, label, val)
        shown += 1
    if shown == 0:
        _add_bullets(slide, ["Breadth data not available in this run."], 1.0, 2.2, 10.8, 2.0, 22)
    _add_bullets(slide, report.get("breadth_notes", [])[:3], 1.0, 4.1, 11.2, 1.9, 16)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "Macro & cross-asset dashboard")
    macro = report.get("macro", []) or []
    if macro:
        for i, row in enumerate(macro[:8]):
            col = i % 4; row_idx = i // 4
            change = row.get("change_pct")
            _metric_card(slide, 0.8 + col * 3.05, 1.75 + row_idx * 1.65, 2.7, 1.25, str(row.get("label", row.get("symbol", ""))), str(row.get("display_value", "Not available")), None if change is None else float(change) >= 0)
    else:
        _add_bullets(slide, ["Cross-asset data not available in this run."], 1.0, 2.2, 10.8, 2.0, 22)
    _add_bullets(slide, report.get("macro_notes", [])[:3], 0.95, 5.15, 11.4, 1.4, 14)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "Technical context")
    technical = report.get("technical", []) or []
    if technical:
        y = 1.7
        for row in technical[:6]:
            _metric_card(slide, 0.9, y, 3.5, 0.78, str(row.get("label", "Metric")), str(row.get("value", "Not available")))
            y += 0.88
        _add_bullets(slide, report.get("technical_notes", [])[:4], 4.8, 1.8, 7.2, 4.5, 16)
    else:
        _add_bullets(slide, ["Technical metrics not available in this run. No levels were inferred."], 1.0, 2.2, 11.0, 2.0, 21)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "What drove the week")
    drivers = report.get("drivers", []) or []
    _add_bullets(slide, [f"{d.get('title')}: {d.get('detail')}" for d in drivers[:5]] or ["No verified drivers supplied."], 0.9, 1.7, 11.4, 4.9, 16)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "Risks / what to watch")
    risks = report.get("risks", []) or []
    _add_bullets(slide, [f"{r.get('title')}: {r.get('detail')}" for r in risks[:5]] or ["No watch items supplied."], 0.9, 1.7, 11.4, 4.9, 17)
    _add_footer(slide, footer)

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _add_title(slide, "Final takeaway")
    final = report.get("final_takeaway", {}) or {}
    cards = [("Market regime", final.get("market_regime", "Not available")), ("Leadership", final.get("leadership", "Not available")), ("Key tension", final.get("key_tension", "Not available"))]
    for i, (label, value) in enumerate(cards):
        _metric_card(slide, 0.9 + i * 4.1, 2.0, 3.7, 1.45, label, str(value))
    summary = final.get("summary", "")
    if summary:
        box = slide.shapes.add_textbox(Inches(1.0), Inches(4.15), Inches(11.2), Inches(1.3))
        p = box.text_frame.paragraphs[0]
        p.text = summary
        p.font.name = "Aptos Display"; p.font.size = Pt(22); p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
    _add_footer(slide, footer)

    prs.save(output)
    return output
